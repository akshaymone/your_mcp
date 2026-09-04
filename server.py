import os
import sys
import logging
import base64
import uuid
import tempfile
import subprocess
import requests
from pathlib import Path
from io import BytesIO

from mcp.server.mcpserver import MCPServer
from dotenv import load_dotenv

# Load environment variables from .env
load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# Initialize MCP Server
mcp = MCPServer("OpenCode-MCP-Extensions")

# ── Infrastructure Pre-reqs ───────────────────────────────────────────────────

def check_qdrant_status():
    """Pings Qdrant to see if it's running. Auto-starts if down."""
    qdrant_host = os.getenv("QDRANT_HOST", "127.0.0.1")
    qdrant_port = os.getenv("QDRANT_PORT", "6333")
    url = f"http://{qdrant_host}:{qdrant_port}"
    try:
        response = requests.get(url, timeout=2.0)
        if response.status_code == 200:
            logger.info("Qdrant is running.")
            return True
    except requests.exceptions.RequestException:
        pass
        
    logger.warning("Qdrant not reachable. Attempting to start via docker compose...")
    try:
        subprocess.run(["docker", "compose", "up", "-d"], check=True, capture_output=True)
        logger.info("Docker compose command successful.")
        return True
    except FileNotFoundError:
        logger.error("Docker command not found.")
        return False
    except subprocess.CalledProcessError as e:
        logger.error(f"Failed to start Qdrant via docker: {e.stderr.decode()}")
        return False

# Attempt to start Qdrant on init
check_qdrant_status()

# ── Tools ─────────────────────────────────────────────────────────────────────

@mcp.tool()
def convert_office_to_pdf(input_path: str, output_dir: str) -> str:
    """
    Uses local win32com to silently convert a DOCX or PPTX file to a PDF.
    
    Args:
        input_path: Absolute path to the Word or PowerPoint file.
        output_dir: Absolute path to the directory where the PDF should be saved.
        
    Returns:
        The absolute path to the generated PDF.
    """
    import pythoncom
    import win32com.client

    input_path = Path(input_path).resolve()
    if not input_path.exists():
        return f"Error: Cannot find file at {input_path}"
        
    ext = input_path.suffix.lower()
    if ext == ".pdf":
        return str(input_path)
        
    output_path = Path(output_dir) / input_path.with_suffix(".pdf").name
    
    pythoncom.CoInitialize()
    try:
        if ext == ".docx":
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            try:
                doc = word.Documents.Open(str(input_path))
                doc.SaveAs(str(output_path), FileFormat=17) # 17 is wdFormatPDF
                doc.Close()
            finally:
                word.Quit()
        elif ext == ".pptx":
            powerpoint = win32com.client.DispatchEx("Powerpoint.Application")
            try:
                presentation = powerpoint.Presentations.Open(str(input_path), WithWindow=False)
                presentation.SaveAs(str(output_path), 32) # 32 is ppSaveAsPDF
                presentation.Close()
            finally:
                powerpoint.Quit()
        else:
            return f"Error: Unsupported file format {ext}. Only .docx and .pptx are supported."
    except Exception as e:
        return f"Error converting document: {e}"
    finally:
        pythoncom.CoUninitialize()
        
    return str(output_path)


@mcp.tool()
def extract_pdf_pages(pdf_path: str, output_dir: str = None) -> str:
    """
    Slices a PDF into a set of JPEG images (one per page) using pdf2image.
    If output_dir is not provided, it writes to a temporary directory.
    
    Args:
        pdf_path: Absolute path to the PDF.
        output_dir: Optional absolute path to save the JPEGs.
        
    Returns:
        A JSON string containing the list of extracted image paths.
    """
    import json
    from pdf2image import convert_from_path, pdfinfo_from_path
    
    try:
        pdf_path = Path(pdf_path)
        if not pdf_path.exists():
            return f"Error: Cannot find PDF at {pdf_path}"
            
        if not output_dir:
            temp_dir = tempfile.mkdtemp(prefix="pdf_extract_")
            output_dir = temp_dir
            
        out_path = Path(output_dir)
        out_path.mkdir(parents=True, exist_ok=True)
        
        info = pdfinfo_from_path(str(pdf_path))
        total_pages = info.get("Pages", 1)
        
        image_paths = []
        batch_size = 5
        
        for start_page in range(1, total_pages + 1, batch_size):
            end_page = min(start_page + batch_size - 1, total_pages)
            images = convert_from_path(str(pdf_path), first_page=start_page, last_page=end_page)
            
            for i, img in enumerate(images):
                page_num = start_page + i
                img_path = out_path / f"page_{page_num:03d}.jpg"
                img.save(str(img_path), "JPEG")
                image_paths.append(str(img_path))
                
        return json.dumps({"output_dir": output_dir, "images": image_paths})
        
    except FileNotFoundError as e:
        if "poppler" in str(e).lower():
            return "Error: Poppler is not installed or not in PATH. Please instruct the user to install Poppler."
        return f"Error: {e}"
    except Exception as e:
        return f"Error extracting pages: {e}"


@mcp.tool()
def index_images_to_qdrant(image_paths: list[str], collection_name: str, doc_name: str) -> str:
    """
    Embeds page images via ColPali and stores them in Qdrant along with payload metadata.
    
    Args:
        image_paths: List of absolute paths to the JPEG page images.
        collection_name: Target Qdrant collection name (e.g., 'vision_pages').
        doc_name: A logical name for the document to group these pages under.
        
    Returns:
        Status message about the indexing.
    """
    import torch
    from PIL import Image
    from colpali_engine.models import ColIdefics3, ColIdefics3Processor
    from qdrant_client import QdrantClient
    from qdrant_client.http.models import Distance, VectorParams, PointStruct, MultiVectorConfig, MultiVectorComparator
    
    qdrant_host = os.getenv("QDRANT_HOST", "127.0.0.1")
    qdrant_port = os.getenv("QDRANT_PORT", "6333")
    
    try:
        qdrant = QdrantClient(host=qdrant_host, port=qdrant_port, timeout=10.0)
    except Exception as e:
        return f"Error: Could not connect to Qdrant at {qdrant_host}:{qdrant_port}. Is Docker running? Error: {e}"
        
    # Ensure collection
    if not qdrant.collection_exists(collection_name):
        qdrant.create_collection(
            collection_name=collection_name,
            vectors_config=VectorParams(
                size=128,
                distance=Distance.COSINE,
                multivector_config=MultiVectorConfig(
                    comparator=MultiVectorComparator.MAX_SIM
                )
            )
        )
        
    # Check if doc already indexed (naive check for now, can be improved with hashes)
    point_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{doc_name}_page_1"))
    existing = qdrant.retrieve(collection_name=collection_name, ids=[point_id])
    if existing:
        return f"Document '{doc_name}' is already indexed in collection '{collection_name}'."
        
    try:
        # Load ColPali locally
        device = "cuda" if torch.cuda.is_available() else "cpu"
        model_name = os.getenv("VISION_RETRIEVER_MODEL")
        if not model_name:
            return "Error: VISION_RETRIEVER_MODEL must be configured in the .env file."
        
        processor = ColIdefics3Processor.from_pretrained(model_name)
        model = ColIdefics3.from_pretrained(model_name, torch_dtype=torch.bfloat16, device_map=device)
        model.eval()
        
        # Load images
        images = [Image.open(p).convert("RGB") for p in image_paths]
        if not images:
            return "Error: No images found."
            
        inputs = processor.process_images(images).to(device)
        with torch.no_grad():
            embeddings = model(**inputs)
            multi_vectors = embeddings.cpu().float().numpy().tolist()
            
        points = []
        for i, (img_path, mv) in enumerate(zip(image_paths, multi_vectors)):
            page_num = i + 1
            point_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{doc_name}_page_{page_num}"))
            
            with open(img_path, "rb") as f:
                img_str = base64.b64encode(f.read()).decode()
                
            points.append(PointStruct(
                id=point_id,
                vector=mv,
                payload={
                    "doc_name": doc_name,
                    "page_number": page_num,
                    "image_base64": img_str,
                    "file_path": img_path
                }
            ))
            
        qdrant.upsert(collection_name=collection_name, points=points)
        return f"Successfully indexed {len(points)} pages for '{doc_name}' into '{collection_name}'."
        
    except Exception as e:
        return f"Error during indexing: {e}"


@mcp.tool()
def search_visual_knowledge_base(query: str = "", collection_name: str = "vision_pages", limit: int = 5, fetch_doc: str = None, fetch_page: int = None) -> str:
    """
    Embeds a text query using ColPali, searches Qdrant, and returns matching page metadata.
    It also supports exact-match fetching if fetch_doc and fetch_page are provided (used for <FETCH_PAGE> behavior).
    
    CRITICAL INSTRUCTIONS FOR THE AGENT:
    When you call this tool, you MUST execute the following Map-Reduce flow:
    1. MAP: Iterate over the retrieved pages. For each result:
       - If 'file_path' is non-null → call analyze_image(file_path, prompt). This is the fast path.
       - If 'file_path' is null (legacy doc) → call analyze_image(image_base64, prompt) instead.
    2. REDUCE: Read all your extractions and synthesize the final answer.
    3. FALLBACK: If a technical term is undefined in the visual context, use your pre-trained knowledge but label it explicitly as [General Knowledge].
    
    Args:
        query: The search text (leave empty if using fetch_doc/fetch_page).
        collection_name: The Qdrant collection to search.
        limit: Max number of pages to return for semantic search.
        fetch_doc: Optional document name to fetch exactly.
        fetch_page: Optional page number to fetch exactly.
        
    Returns:
        JSON string containing search results with doc_name, page_number, score,
        and either 'file_path' (preferred, pass directly to analyze_image) or
        'image_base64' as fallback for legacy documents indexed before this fix.
    """
    import json
    import torch
    import uuid
    from colpali_engine.models import ColIdefics3, ColIdefics3Processor
    from qdrant_client import QdrantClient
    
    qdrant_host = os.getenv("QDRANT_HOST", "127.0.0.1")
    qdrant_port = os.getenv("QDRANT_PORT", "6333")
    
    try:
        qdrant = QdrantClient(host=qdrant_host, port=qdrant_port, timeout=10.0)
    except Exception as e:
        return f"Error connecting to Qdrant: {e}"
        
    if not qdrant.collection_exists(collection_name):
        return f"Error: Collection '{collection_name}' does not exist."
        
    try:
        out = []
        
        if fetch_doc and fetch_page:
            # Exact match fetching
            point_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{fetch_doc}_page_{fetch_page}"))
            results = qdrant.retrieve(collection_name=collection_name, ids=[point_id])
            for r in results:
                file_path = r.payload.get("file_path")
                out.append({
                    "score": 1.0,
                    "doc_name": r.payload.get("doc_name"),
                    "page_number": r.payload.get("page_number"),
                    # Prefer stable file_path (pass directly to analyze_image).
                    # Falls back to full base64 only when file_path is absent
                    # (docs indexed before this fix).
                    "file_path": file_path if file_path and os.path.exists(file_path) else None,
                    "image_base64": None if (file_path and os.path.exists(file_path)) else r.payload.get("image_base64", ""),
                })
        elif query:
            # Semantic Search
            device = "cuda" if torch.cuda.is_available() else "cpu"
            model_name = os.getenv("VISION_RETRIEVER_MODEL")
            if not model_name:
                return "Error: VISION_RETRIEVER_MODEL must be configured in the .env file."
            
            processor = ColIdefics3Processor.from_pretrained(model_name)
            model = ColIdefics3.from_pretrained(model_name, torch_dtype=torch.bfloat16, device_map=device)
            model.eval()
            
            inputs = processor.process_queries([query]).to(device)
            with torch.no_grad():
                embeddings = model(**inputs)
                query_vector = embeddings[0].cpu().float().numpy().tolist()
                
            results = qdrant.query_points(
                collection_name=collection_name,
                query=query_vector,
                limit=limit
            )
            
            for r in results.points:
                file_path = r.payload.get("file_path")
                out.append({
                    "score": r.score,
                    "doc_name": r.payload.get("doc_name"),
                    "page_number": r.payload.get("page_number"),
                    # Prefer stable file_path (pass directly to analyze_image).
                    # Falls back to full base64 only when file_path is absent
                    # (docs indexed before this fix).
                    "file_path": file_path if file_path and os.path.exists(file_path) else None,
                    "image_base64": None if (file_path and os.path.exists(file_path)) else r.payload.get("image_base64", ""),
                })
        else:
            return "Error: Must provide either a search 'query' or both 'fetch_doc' and 'fetch_page'."
            
        return json.dumps({"query": query, "results": out})
        
    except Exception as e:
        return f"Error searching: {e}"


@mcp.tool()
def get_file_info(file_path: str) -> str:
    """
    Returns basic metadata about any file on disk without ingesting it.

    Use this when the user asks things like:
      - "can you access this file path?"
      - "what is this file?"
      - "what's the file size / format / author of <path>?"
      - "how many pages does this PDF have?"

    Args:
        file_path: Absolute path to any file (PDF, PPTX, DOCX, etc.).

    Returns:
        A JSON string with file metadata: name, size, extension, dates,
        and for PDFs: page count, title, author, subject from PDF metadata.
    """
    import os
    import json
    from pathlib import Path
    from datetime import datetime

    path = Path(file_path.strip().strip('"').strip("'"))

    if not path.exists():
        return json.dumps({"error": f"File not found: {file_path}"})
    if not path.is_file():
        return json.dumps({"error": f"Path is not a file: {file_path}"})

    stat = path.stat()
    info = {
        "file_name":   path.name,
        "extension":   path.suffix.lower(),
        "size_bytes":  stat.st_size,
        "size_kb":     round(stat.st_size / 1024, 2),
        "size_mb":     round(stat.st_size / (1024 * 1024), 3),
        "created":     datetime.fromtimestamp(stat.st_ctime).isoformat(),
        "modified":    datetime.fromtimestamp(stat.st_mtime).isoformat(),
        "full_path":   str(path.resolve()),
    }

    # PDF-specific metadata
    if path.suffix.lower() == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            meta = reader.metadata or {}
            info["pdf_pages"]   = len(reader.pages)
            info["pdf_title"]   = meta.get("/Title", "")
            info["pdf_author"]  = meta.get("/Author", "")
            info["pdf_subject"] = meta.get("/Subject", "")
            info["pdf_creator"] = meta.get("/Creator", "")
        except ImportError:
            info["pdf_note"] = "pypdf not installed; PDF metadata unavailable."
        except Exception as e:
            info["pdf_note"] = f"Could not read PDF metadata: {e}"

    # PPTX-specific metadata
    elif path.suffix.lower() == ".pptx":
        try:
            from pptx import Presentation
            prs = Presentation(str(path))
            core = prs.core_properties
            info["pptx_slides"]   = len(prs.slides)
            info["pptx_title"]    = core.title or ""
            info["pptx_author"]   = core.author or ""
            info["pptx_subject"]  = core.subject or ""
            info["pptx_modified"] = core.modified.isoformat() if core.modified else ""
        except Exception as e:
            info["pptx_note"] = f"Could not read PPTX metadata: {e}"

    # DOCX-specific metadata
    elif path.suffix.lower() == ".docx":
        try:
            import docx
            doc = docx.Document(str(path))
            core = doc.core_properties
            info["docx_title"]    = core.title or ""
            info["docx_author"]   = core.author or ""
            info["docx_subject"]  = core.subject or ""
            info["docx_modified"] = core.modified.isoformat() if core.modified else ""
        except Exception as e:
            info["docx_note"] = f"Could not read DOCX metadata: {e}"

    return json.dumps(info, ensure_ascii=False, indent=2)


@mcp.tool()
def analyze_image(image_path_or_base64: str, prompt: str) -> str:
    """
    Passes an image (either an absolute path or a base64 string) and a text prompt to the Vision LLM to extract data.
    
    Args:
        image_path_or_base64: Path to local image OR the base64 encoded string from search results.
        prompt: Specific question or extraction instruction for the model.
        
    Returns:
        The extracted text or analysis from the Vision LLM.
    """
    import os
    import base64
    import requests
    import urllib3

    # The internal FM Gateway uses a self-signed / corporate CA certificate.
    # Disable SSL verification and suppress the resulting InsecureRequestWarning.
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    api_url = os.getenv("FM_GATEWAY_URL")
    api_key = os.getenv("FM_GATEWAY_TOKEN", "")
    model_name = os.getenv("VLM_MODEL")
    
    if not api_url or not model_name:
        return "Error: FM_GATEWAY_URL and VLM_MODEL must be configured in the .env file."
    
    # Normalize URL: strip trailing slashes and any existing /v1 suffix,
    # then always append /v1/chat/completions for the OpenAI-compatible endpoint.
    base_url = api_url.rstrip("/")
    if base_url.endswith("/v1"):
        base_url = base_url[:-3]  # strip /v1 so we always add it cleanly below
    endpoint = f"{base_url}/v1/chat/completions"
    logger.info(f"analyze_image: calling endpoint={endpoint} model={model_name}")

    # Handle local file paths vs base64 strings
    if os.path.exists(image_path_or_base64):
        logger.info(f"analyze_image: reading image from file: {image_path_or_base64}")
        try:
            with open(image_path_or_base64, "rb") as f:
                base64_img = base64.b64encode(f.read()).decode("utf-8")
        except OSError as e:
            err_msg = f"Error: Could not read image file '{image_path_or_base64}': {e}"
            logger.error(err_msg)
            return err_msg
    else:
        # Validate it looks like a real base64 string (not a truncated display artifact)
        candidate = image_path_or_base64.strip()
        if not candidate or "truncated" in candidate.lower() or len(candidate) < 100:
            err_msg = (
                f"Error: 'image_path_or_base64' does not point to an existing file and does not "
                f"look like valid base64 data (length={len(candidate)}). "
                f"Pass the 'file_path' returned by search_visual_knowledge_base instead."
            )
            logger.error(err_msg)
            return err_msg
        logger.info("analyze_image: using raw base64 string as image input.")
        base64_img = candidate

    headers = {
        "Content-Type": "application/json"
    }
    if api_key:
        headers["Authorization"] = f"Bearer {api_key}"

    payload = {
        "model": model_name,
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:image/jpeg;base64,{base64_img}"
                        }
                    }
                ]
            }
        ],
        "max_tokens": 1500
    }

    try:
        response = requests.post(endpoint, headers=headers, json=payload, timeout=60, verify=False)
        response.raise_for_status()
    except requests.exceptions.HTTPError as e:
        err_msg = (
            f"API HTTP Error {e.response.status_code} from FM Gateway.\n"
            f"Endpoint: {endpoint}\nModel: {model_name}\nResponse Body: {e.response.text}"
        )
        logger.error(err_msg)
        return err_msg
    except requests.exceptions.ConnectionError as e:
        err_msg = f"Connection Error — could not reach FM Gateway at {endpoint}: {e}"
        logger.error(err_msg)
        return err_msg
    except requests.exceptions.Timeout:
        err_msg = f"Timeout Error — FM Gateway at {endpoint} did not respond within 60s."
        logger.error(err_msg)
        return err_msg
    except requests.exceptions.RequestException as e:
        err_msg = f"Request Error calling FM Gateway: {e}"
        logger.error(err_msg)
        return err_msg

    # Parse the JSON response — gateway may return a non-JSON body on some errors
    try:
        result = response.json()
    except Exception as e:
        err_msg = (
            f"Error: FM Gateway returned non-JSON response (status {response.status_code}).\n"
            f"Raw body: {response.text[:500]}"
        )
        logger.error(err_msg)
        return err_msg

    # Extract content — guard against unexpected response shapes
    try:
        return result["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError) as e:
        err_msg = (
            f"Error: Unexpected response structure from FM Gateway: {e}\n"
            f"Full response: {result}"
        )
        logger.error(err_msg)
        return err_msg


@mcp.tool()
def delete_document(doc_name_or_path: str, collection_name: str = "vision_pages") -> str:
    """
    Deletes a specific document — or ALL documents — from the visual knowledge base.

    Use this when the user says things like:
      - "delete report from my knowledge base"
      - "remove presentation.pptx from the knowledge base"
      - "delete all documents from my local knowledge base"
      - "wipe / clear the knowledge base"

    Args:
        doc_name_or_path: The document to delete. Accepts:
            - A plain doc name:  "report"
            - A filename:        "report.pptx"
            - A full file path:  "C:\\path\\to\\report.pptx"
            - The keyword "all", "*", or "everything" to wipe ALL documents.
        collection_name: Target Qdrant collection (default: 'vision_pages').

    Returns:
        Confirmation message with the count of deleted pages, or an error description.
    """
    from qdrant_client import QdrantClient
    from qdrant_client.http.models import Filter, FieldCondition, MatchValue

    qdrant_host = os.getenv("QDRANT_HOST", "127.0.0.1")
    qdrant_port = os.getenv("QDRANT_PORT", "6333")

    try:
        qdrant = QdrantClient(host=qdrant_host, port=qdrant_port, timeout=10.0)
    except Exception as e:
        return f"Error: Could not connect to Qdrant: {e}"

    if not qdrant.collection_exists(collection_name):
        return f"Collection '{collection_name}' does not exist. Nothing to delete."

    # ── Delete ALL documents ──────────────────────────────────────────────────
    if doc_name_or_path.strip().lower() in ("all", "*", "everything"):
        # Count total points before wiping
        collection_info = qdrant.get_collection(collection_name)
        total = collection_info.points_count
        qdrant.delete_collection(collection_name)
        return (
            f"✅ Cleared entire knowledge base: deleted all {total} page(s) "
            f"and removed collection '{collection_name}'. "
            f"It will be recreated automatically on the next ingestion."
        )

    # ── Delete a specific document ────────────────────────────────────────────
    # Derive doc_name from path, filename, or bare name — always via Path.stem
    doc_name = Path(doc_name_or_path.strip()).stem

    # Count pages for this doc before deleting
    scroll_result, _ = qdrant.scroll(
        collection_name=collection_name,
        scroll_filter=Filter(must=[FieldCondition(key="doc_name", match=MatchValue(value=doc_name))]),
        limit=1000,
        with_payload=False,
        with_vectors=False,
    )

    if not scroll_result:
        return (
            f"No document named '{doc_name}' found in collection '{collection_name}'. "
            f"Nothing was deleted."
        )

    page_count = len(scroll_result)

    qdrant.delete(
        collection_name=collection_name,
        points_selector=Filter(must=[FieldCondition(key="doc_name", match=MatchValue(value=doc_name))]),
    )

    return f"✅ Deleted '{doc_name}' ({page_count} page(s)) from collection '{collection_name}'."


@mcp.tool()
def check_document_status(file_path: str, collection_name: str = "vision_pages") -> str:
    """
    Checks whether a given document has already been indexed in the visual knowledge base.

    IMPORTANT: Call this tool FIRST whenever the user references a document file path alongside
    any question or request. Use the result to decide whether to ingest first or query directly.

    Args:
        file_path: Absolute path to the document file (e.g., C:\\path\\to\\report.pptx).
        collection_name: The Qdrant collection to check (default: 'vision_pages').

    Returns:
        JSON string with 'status' of "indexed" (includes page_count) or "not_indexed", plus the
        derived doc_name that will be used for all subsequent ingestion and retrieval operations.
    """
    import json
    from qdrant_client import QdrantClient
    from qdrant_client.http.models import Filter, FieldCondition, MatchValue

    doc_name = Path(file_path).stem
    qdrant_host = os.getenv("QDRANT_HOST", "127.0.0.1")
    qdrant_port = os.getenv("QDRANT_PORT", "6333")

    try:
        qdrant = QdrantClient(host=qdrant_host, port=qdrant_port, timeout=5.0)
    except Exception as e:
        return json.dumps({"status": "error", "message": f"Could not connect to Qdrant: {e}"})

    if not qdrant.collection_exists(collection_name):
        return json.dumps({"status": "not_indexed", "doc_name": doc_name, "reason": "collection does not exist yet"})

    # Check for page_1 as a reliable proxy for whether the document is indexed
    point_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{doc_name}_page_1"))
    existing = qdrant.retrieve(collection_name=collection_name, ids=[point_id], with_payload=False)

    if existing:
        # Count total pages indexed for this document
        scroll_result, _ = qdrant.scroll(
            collection_name=collection_name,
            scroll_filter=Filter(must=[FieldCondition(key="doc_name", match=MatchValue(value=doc_name))]),
            limit=1000,
            with_payload=False,
            with_vectors=False,
        )
        page_count = len(scroll_result)
        return json.dumps({
            "status": "indexed",
            "doc_name": doc_name,
            "page_count": page_count,
            "collection": collection_name,
        })
    else:
        return json.dumps({
            "status": "not_indexed",
            "doc_name": doc_name,
            "collection": collection_name,
        })


@mcp.tool()
def ingest_document(file_path: str, collection_name: str = "vision_pages") -> str:
    """
    High-level end-to-end ingestion tool. Converts a document to PDF, extracts pages as images,
    and indexes them into the visual knowledge base using ColPali embeddings.

    Use this tool when check_document_status reports 'not_indexed'. Before calling this tool,
    tell the user: "This document hasn't been ingested yet — please wait while I process it."

    Args:
        file_path: Absolute path to the document (.pptx, .docx, or .pdf).
        collection_name: Target Qdrant collection (default: 'vision_pages').

    Returns:
        A step-by-step progress log and final status confirming the document is ready to query.
    """
    import json
    import tempfile

    progress = []
    file_path = str(Path(file_path).resolve())
    doc_name = Path(file_path).stem
    ext = Path(file_path).suffix.lower()

    # ── Step 1: Convert to PDF ────────────────────────────────────────────────
    if ext in (".docx", ".pptx"):
        progress.append(f"[1/3] Converting '{Path(file_path).name}' to PDF...")
        temp_pdf_dir = tempfile.mkdtemp(prefix="ingest_pdf_")
        pdf_path = convert_office_to_pdf(file_path, temp_pdf_dir)
        if pdf_path.startswith("Error"):
            return "\n".join(progress) + f"\n❌ Conversion failed: {pdf_path}"
        progress.append(f"[1/3] ✅ PDF saved to: {pdf_path}")
    elif ext == ".pdf":
        pdf_path = file_path
        progress.append(f"[1/3] ✅ File is already a PDF — skipping conversion.")
    else:
        return f"❌ Unsupported file type '{ext}'. Supported formats: .pptx, .docx, .pdf"

    # ── Step 2: Extract pages as images ──────────────────────────────────────
    progress.append(f"[2/3] Extracting individual pages from PDF...")
    extract_result = extract_pdf_pages(pdf_path)
    if extract_result.startswith("Error"):
        return "\n".join(progress) + f"\n❌ Page extraction failed: {extract_result}"

    extract_data = json.loads(extract_result)
    image_paths = extract_data["images"]
    progress.append(f"[2/3] ✅ Extracted {len(image_paths)} pages.")

    # ── Step 3: Embed and index into Qdrant ───────────────────────────────────
    progress.append(
        f"[3/3] Embedding and indexing {len(image_paths)} pages into '{collection_name}' "
        f"(this may take a few minutes while the visual model processes the images)..."
    )
    index_result = index_images_to_qdrant(image_paths, collection_name, doc_name)
    if index_result.startswith("Error"):
        return "\n".join(progress) + f"\n❌ Indexing failed: {index_result}"

    progress.append(f"[3/3] ✅ {index_result}")
    progress.append(f"\n✅ '{doc_name}' is now fully indexed and ready to query!")
    return "\n".join(progress)


@mcp.tool()
def list_ingested_documents(collection_name: str = "vision_pages") -> str:
    """
    Lists all documents currently indexed in the visual knowledge base, along with
    their page counts and the collection they belong to.

    Use this when the user asks things like:
      - "what documents are in my knowledge base?"
      - "what have you already ingested?"
      - "show me all indexed documents"
      - "list everything in the knowledge base"

    Args:
        collection_name: The Qdrant collection to inspect (default: 'vision_pages').

    Returns:
        A JSON string with a list of documents, each containing doc_name, page_count,
        and collection_name. Returns an appropriate message if the collection is empty
        or does not exist.
    """
    import json
    from qdrant_client import QdrantClient

    qdrant_host = os.getenv("QDRANT_HOST", "127.0.0.1")
    qdrant_port = os.getenv("QDRANT_PORT", "6333")

    try:
        qdrant = QdrantClient(host=qdrant_host, port=qdrant_port, timeout=10.0)
    except Exception as e:
        return json.dumps({"error": f"Could not connect to Qdrant: {e}"})

    if not qdrant.collection_exists(collection_name):
        return json.dumps({
            "collection": collection_name,
            "documents": [],
            "message": f"Collection '{collection_name}' does not exist yet. No documents have been ingested."
        })

    try:
        # Scroll through all points and aggregate by doc_name
        all_points = []
        offset = None

        while True:
            batch, next_offset = qdrant.scroll(
                collection_name=collection_name,
                limit=250,
                offset=offset,
                with_payload=["doc_name"],
                with_vectors=False,
            )
            all_points.extend(batch)
            if next_offset is None:
                break
            offset = next_offset

        if not all_points:
            return json.dumps({
                "collection": collection_name,
                "documents": [],
                "message": "The knowledge base is empty — no documents have been ingested yet."
            })

        # Aggregate page counts per document
        doc_counts: dict[str, int] = {}
        for point in all_points:
            name = point.payload.get("doc_name", "unknown")
            doc_counts[name] = doc_counts.get(name, 0) + 1

        documents = [
            {"doc_name": name, "page_count": count, "collection": collection_name}
            for name, count in sorted(doc_counts.items())
        ]

        return json.dumps({
            "collection": collection_name,
            "total_documents": len(documents),
            "total_pages": sum(d["page_count"] for d in documents),
            "documents": documents
        })
    except Exception as e:
        return json.dumps({"error": f"Error listing documents: {e}"})


@mcp.tool()
def launch_ingestion(ask_me_dir: str, doc_name: str = None) -> str:
    """
    Launches 'ask-me ingest' as a fire-and-forget background subprocess.
    Returns immediately — no MCP timeout risk regardless of document size.

    Use this instead of ingest_document for any document ingestion.
    After calling this tool, monitor progress using poll_ingestion_status().

    Workflow:
        1. Call launch_ingestion(ask_me_dir='/path/to/ask-me') — returns immediately.
        2. Tell the user: "Ingestion started in the background. I'll monitor progress."
        3. Poll poll_ingestion_status(doc_name) every 30s until status == 'done' or 'error'.
        4. Once done, proceed to query with search_visual_knowledge_base.

    Args:
        ask_me_dir: Absolute path to the ask-me project directory (contains pyproject.toml).
        doc_name: Optional document name for the user-facing message. Does not affect the process.

    Returns:
        Confirmation that the background process was launched, with its PID.
    """
    import subprocess
    import sys

    ask_me_path = Path(ask_me_dir).resolve()
    if not ask_me_path.exists():
        return f"Error: ask-me directory not found at '{ask_me_dir}'."

    python = sys.executable

    try:
        proc = subprocess.Popen(
            [python, "-m", "ask_me.main", "ingest"],
            cwd=str(ask_me_path),
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            # Detach fully from the MCP server process so it keeps running after return
            start_new_session=True,
        )
        msg = (
            f"✅ ask-me ingestion launched as background process (PID {proc.pid}). "
            f"{'Document: ' + doc_name + '.' if doc_name else ''} "
            f"Use poll_ingestion_status(doc_name) to monitor progress."
        )
        logger.info(msg)
        return msg
    except Exception as e:
        err = f"Error launching ask-me: {e}"
        logger.error(err)
        return err


@mcp.tool()
def poll_ingestion_status(doc_name: str) -> str:
    """
    Reads the status file written by ask-me during ingestion and returns current progress.

    Call this every 30 seconds after launch_ingestion() until status is 'done' or 'error'.

    Status values:
        - 'running'  → still in progress; pages_done shows how many pages are embedded so far.
        - 'done'     → ingestion complete; document is ready to query.
        - 'error'    → ingestion failed; 'error' field contains the reason.
        - 'not_found'→ no status file exists yet (process hasn't started or doc_name is wrong).

    Args:
        doc_name: The document stem name (e.g. 'report' for 'report.pptx').

    Returns:
        JSON string with status, pages_done, pages_total, finished_at, and error fields.
    """
    import json

    status_dir = Path(os.path.expanduser("~")) / ".ask_me_store" / "status"
    # Also check STATUS_DIR from env if set
    env_status_dir = os.getenv("STATUS_DIR")
    if env_status_dir:
        status_dir = Path(env_status_dir)

    status_file = status_dir / f"{doc_name}.json"

    if not status_file.exists():
        return json.dumps({
            "doc_name": doc_name,
            "status": "not_found",
            "message": (
                f"No status file found for '{doc_name}'. "
                f"Either the ingestion hasn't started yet or the doc_name is incorrect. "
                f"Expected file: {status_file}"
            )
        })

    try:
        data = json.loads(status_file.read_text(encoding="utf-8"))
        return json.dumps(data, indent=2)
    except Exception as e:
        return json.dumps({"error": f"Could not read status file: {e}"})


if __name__ == "__main__":
    mcp.run(transport="sse", port=8000)
